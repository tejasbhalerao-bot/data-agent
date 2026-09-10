"""
Q1: avg(DATE(digitised_delivery_promise) - DATE(digitised_ts)) across all orders
    with both timestamps present.
Q2: % of orders where pickup_time IS NOT NULL AND
    DATE(delivery_attempt_time) = DATE(digitised_delivery_promise).
    Denominator = all orders in the file.
"""

import warnings
from datetime import datetime

import openpyxl

warnings.filterwarnings("ignore")

XLSX = 'archives/egregiously-miscalibrated-promises/raw-data/all-orders-august-2026.xlsx'

wb = openpyxl.load_workbook(XLSX, read_only=True)
ws = wb[wb.sheetnames[0]]

rows = ws.iter_rows(values_only=True)
header = next(rows)
idx = {name: i for i, name in enumerate(header)}

col_digitised_ts = idx['digitised_ts']
col_delivery_promise = idx['digitised_delivery_promise']
col_pickup_time = idx['pickup_time']
col_delivery_attempt = idx['delivery_attempt_time']


def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if hasattr(v, 'year'):
        return v
    s = str(v).strip()
    if not s:
        return None
    try:
        return datetime.fromisoformat(s[:19]).date()
    except Exception:
        return None


total_orders = 0
digitised_ts_min = None
digitised_ts_max = None

q1_n = 0
q1_sum_days = 0

q2_numerator = 0

for row in rows:
    total_orders += 1

    dig_ts = row[col_digitised_ts]
    if isinstance(dig_ts, datetime):
        if digitised_ts_min is None or dig_ts < digitised_ts_min:
            digitised_ts_min = dig_ts
        if digitised_ts_max is None or dig_ts > digitised_ts_max:
            digitised_ts_max = dig_ts

    dig_ts_date = to_date(dig_ts)
    delivery_promise_date = to_date(row[col_delivery_promise])
    if dig_ts_date and delivery_promise_date:
        q1_n += 1
        q1_sum_days += (delivery_promise_date - dig_ts_date).days

    pickup = row[col_pickup_time]
    delivery_attempt_date = to_date(row[col_delivery_attempt])
    if pickup is not None and delivery_attempt_date and delivery_promise_date:
        if delivery_attempt_date == delivery_promise_date:
            q2_numerator += 1

print(f"Total orders in file           : {total_orders:,}")
print(f"digitised_ts range             : {digitised_ts_min} -> {digitised_ts_max}")
print()

print("=" * 74)
print("Q1 — avg(DATE(digitised_delivery_promise) - DATE(digitised_ts))")
print("=" * 74)
print(f"Orders with both timestamps    : {q1_n:,}")
print(f"Average lead time (days)       : {q1_sum_days / q1_n:.4f}")
print()

print("=" * 74)
print("Q2 — % pickup_time NOT NULL AND DATE(delivery_attempt_time) = DATE(digitised_delivery_promise)")
print("=" * 74)
print(f"Numerator (matching orders)    : {q2_numerator:,}")
print(f"Denominator (all orders)       : {total_orders:,}")
print(f"Percentage                     : {100 * q2_numerator / total_orders:.4f}%")
